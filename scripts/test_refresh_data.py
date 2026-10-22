"""Offline integrity tests for a public comparison built from mutable statistics."""

import copy
import datetime as dt
import importlib.util
import json
from pathlib import Path
import unittest
from unittest.mock import patch

import openpyxl

SCRIPT = Path(__file__).with_name("refresh_data.py")
SPEC = importlib.util.spec_from_file_location("refresh_data", SCRIPT)
refresh = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(refresh)


class SourceTransformTests(unittest.TestCase):
    def test_financial_quarter_changes_calendar_year(self):
        self.assertEqual(refresh.quarter_end("2025/26", 1), dt.date(2025, 6, 30))
        self.assertEqual(refresh.quarter_end("2025/26", 4), dt.date(2026, 3, 31))

    def test_fraud_does_not_enter_geographical_crime_totals(self):
        self.assertEqual(refresh.categories_for_offence("Fraud offences", "Fraud: Action Fraud"), [])
        self.assertEqual(refresh.categories_for_offence("Theft offences", "Shoplifting"), ["all-crime", "theft", "shoplifting"])

    def test_burglary_history_handles_offence_reclassification(self):
        for subgroup in ("Domestic burglary", "Non-domestic burglary", "Residential burglary", "Non-residential burglary"):
            child = "non-residential-burglary" if subgroup.startswith("Non-") else "residential-burglary"
            self.assertEqual(refresh.categories_for_offence("Theft offences", subgroup), ["all-crime", "theft", "burglary", child])

    def test_violence_components_are_not_counted_as_separate_main_groups(self):
        self.assertEqual(refresh.categories_for_offence("Violence against the person", "Homicide"), ["all-crime", "violence", "homicide"])
        self.assertEqual(refresh.categories_for_offence("Violence against the person", "Stalking and harassment"), ["all-crime", "violence", "stalking-harassment"])

    def test_source_percentage_suppression_is_preserved(self):
        change, reason = refresh.published_change("[u]", 15, 3)
        self.assertIsNone(change)
        self.assertIn("fewer than 50", reason)
        # Other collections sometimes publish a percentage for small counts.
        # Preserve that publication decision rather than imposing a new rule.
        self.assertEqual(refresh.published_change(100, 4, 2), (100, None))
        with self.assertRaisesRegex(ValueError, "disagrees"):
            refresh.published_change(30, 100, 80)

    def test_weapon_date_headers_do_not_mistake_percentages_for_counts(self):
        self.assertEqual(refresh.header_period("Jul 2025 to\n Jun 2026"), (dt.date(2025, 7, 1), dt.date(2026, 6, 30)))
        self.assertIsNone(refresh.header_period("% involving a knife Apr 2023 to Mar 2024"))
        self.assertIsNone(refresh.header_period("Jul 2025 to Jun 2026 compared with previous year % change"))

    def weapon_workbook(self):
        workbook = openpyxl.Workbook()
        workbook.remove(workbook.active)
        for title, headers, values in [
            ("Table P5", ["Total of selected serious offences involving a knife"], [100]),
            ("Table P6", ["Apr 2023 to Mar 2024", "% involving a knife Apr 2023 to Mar 2024", "Jul 2024 to Jun 2025", "Jul 2025 to Jun 2026", "Jul 2025 to Jun 2026 compared with previous year % change"], [70, 0.3, 80, 100, 25]),
            ("Table P7", ["Jul 2024 to Jun 2025", "Jul 2025 to Jun 2026"], [8, 10]),
            ("Table P8", ["Apr 2023 to Mar 2024", "Jul 2024 to Jun 2025", "Jul 2025 to Jun 2026", "Jul 2025 to Jun 2026 compared with previous year % change"], [20, 30, 45, 50]),
        ]:
            sheet = workbook.create_sheet(title)
            sheet.append(["Area Code", "Area Name", *headers])
            sheet.append(["E23000013", "Cleveland", *values])
        return workbook

    def test_weapon_history_retains_real_irregular_dates(self):
        weapons = refresh.parse_weapons(self.weapon_workbook(), dt.date(2026, 6, 30))
        knife = weapons["knife-enabled"]["E23000013"]
        self.assertEqual(list(knife["dated"]), [dt.date(2024, 3, 31), dt.date(2025, 6, 30), dt.date(2026, 6, 30)])
        self.assertEqual(knife["count"], 100)
        self.assertEqual(knife["previousCount"], 80)
        self.assertEqual(knife["changePercent"], 25)
        self.assertEqual(weapons["firearms"]["E23000013"]["count"], 45)

    def test_disagreeing_knife_tables_fail_before_publication(self):
        workbook = self.weapon_workbook()
        workbook["Table P5"].cell(2, 3, 101)
        with self.assertRaisesRegex(ValueError, "P5 and P6"):
            refresh.parse_weapons(workbook, dt.date(2026, 6, 30))

    def test_no_population_is_invented_for_early_weapon_history(self):
        population = {"E23000013": {2011: 100000}}
        self.assertIsNone(refresh.historical_rate("E23000013", dt.date(2011, 3, 31), 70, population, {}))
        self.assertEqual(refresh.historical_rate("E23000013", dt.date(2012, 3, 31), 70, population, {}), 0.7)

    def test_unknown_population_or_stale_denominator_fails(self):
        with self.assertRaisesRegex(ValueError, "No actual"):
            refresh.population_for("E23000099", 2025, {}, {})
        with self.assertRaisesRegex(ValueError, "out of date"):
            refresh.population_for("E23000001", 2025, {"E23000001": {2020: 10}}, {})


class PublishedIntegrityTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.data = json.loads((refresh.OUTPUT / "crime.json").read_text())

    def test_actual_snapshot_has_complete_coverage_and_reconciles(self):
        refresh.validate_dataset(self.data)

    def test_expanded_data_retains_suppression_and_consistent_stored_rate_units(self):
        area = next(a for a in self.data["areas"] if a["id"] == "E23000013")
        self.assertEqual(len(self.data["categories"]), 25)
        suppressed = [m for a in self.data["areas"] for m in a["metrics"].values() if "changeSuppressedReason" in m]
        self.assertTrue(suppressed)
        self.assertTrue(all(m["changePercent"] is None for m in suppressed))
        for cat in refresh.RARE_IDS:
            self.assertAlmostEqual(area["metrics"][cat]["rate"], area["metrics"][cat]["count"] / area["population"] * 1000, places=4)

    def test_national_collections_have_their_published_btp_scope(self):
        national = next(a for a in self.data["areas"] if a["id"] == refresh.NATIONAL)
        forces = [a for a in self.data["areas"] if a["type"] == "force"]
        metadata = {c["id"]: c for c in self.data["categories"]}
        self.assertFalse(metadata["firearms"]["nationalIncludesBtp"])
        self.assertTrue(metadata["knife-enabled"]["nationalIncludesBtp"])
        for field in ("count", "previousCount"):
            self.assertEqual(national["metrics"]["firearms"][field], sum(a["metrics"]["firearms"][field] for a in forces))
            self.assertGreaterEqual(national["metrics"]["knife-enabled"][field], sum(a["metrics"]["knife-enabled"][field] for a in forces))

    def test_suppressed_change_cannot_be_silently_reconstructed(self):
        bad = copy.deepcopy(self.data)
        metric = next(m for a in bad["areas"] for m in a["metrics"].values() if "changeSuppressedReason" in m)
        metric["changePercent"] = 400
        with self.assertRaisesRegex(ValueError, "suppression"):
            refresh.validate_dataset(bad)

    def test_including_btp_in_firearm_baseline_is_rejected(self):
        bad = copy.deepcopy(self.data)
        cat = next(c for c in bad["categories"] if c["id"] == "firearms")
        cat["nationalIncludesBtp"] = True
        with self.assertRaisesRegex(ValueError, "coverage metadata"):
            refresh.validate_dataset(bad)

    def test_missing_force_is_not_treated_as_zero(self):
        bad = copy.deepcopy(self.data)
        bad["areas"] = [a for a in bad["areas"] if a["id"] != "E23000005"]
        with self.assertRaises(ValueError):
            refresh.validate_dataset(bad)

    def test_incorrect_denominator_cannot_publish(self):
        bad = copy.deepcopy(self.data)
        area = next(a for a in bad["areas"] if a["id"] == "E23000005")
        area["population"] *= 2
        with self.assertRaisesRegex(ValueError, "population-normalised"):
            refresh.validate_dataset(bad)

    def test_city_of_london_cannot_enter_rate_rankings(self):
        bad = copy.deepcopy(self.data)
        city = next(a for a in bad["areas"] if a["id"] == refresh.CITY)
        city["comparisonEligible"] = True
        with self.assertRaisesRegex(ValueError, "City of London"):
            refresh.validate_dataset(bad)

    def test_history_endpoint_must_match_latest_count(self):
        bad = copy.deepcopy(self.data)
        bad["areas"][0]["metrics"]["all-crime"]["history"][-1]["count"] += 1
        with self.assertRaisesRegex(ValueError, "Trend endpoint"):
            refresh.validate_dataset(bad)

    def test_failed_fetch_preserves_both_public_snapshots(self):
        before = {filename: (refresh.OUTPUT / filename).read_bytes()
                  for filename in ("crime.json", "source-manifest.json")}
        with patch.object(refresh, "discover_sources", side_effect=RuntimeError("source unavailable")):
            with self.assertRaisesRegex(RuntimeError, "source unavailable"):
                refresh.main([])
        for filename, content in before.items():
            self.assertEqual((refresh.OUTPUT / filename).read_bytes(), content)


if __name__ == "__main__":
    unittest.main()
