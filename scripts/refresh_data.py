#!/usr/bin/env python3
"""Refresh the static England/Wales comparison dataset from official publications.

No API keys or paid services. A failed download, schema change, population gap, or
reconciliation leaves the published JSON files untouched. Run --check to validate
the committed snapshot without network access. The monthly workflow calls this
script; the underlying comparable crime statistics are published quarterly.
"""

from __future__ import annotations

import argparse
import calendar
import collections
import datetime as dt
import hashlib
from html.parser import HTMLParser
import io
import json
import math
import os
from pathlib import Path
import re
import sys
import tempfile
import time
import urllib.error
import urllib.parse
import urllib.request
import xml.etree.ElementTree as ET
import zipfile

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "public/data"
CACHE = ROOT / ".cache/crime-sources"
ONS_INDEX = "https://www.ons.gov.uk/peoplepopulationandcommunity/crimeandjustice/datasets/policeforceareadatatables"
HO_INDEX = "https://www.gov.uk/government/statistical-data-sets/police-recorded-crime-and-outcomes-open-data-tables"
POP_SEARCH = "https://www.ons.gov.uk/search?q=population+estimates+police+force+areas&sortBy=release_date&order=desc"
NATIONAL = "K04000001"
CITY = "E23000034"
CLASSIFICATION_GUIDE = "https://www.ons.gov.uk/peoplepopulationandcommunity/crimeandjustice/methodologies/userguidetocrimestatisticsforenglandandwales#offence-types"
ROBBERY_RECORDING_NOTE = "An April 2025 clarification of police counting rules increased the recording of business robbery during shop theft involving violence or threats. Changes in robbery and shoplifting across this date can reflect the recording change as well as changes in crime."
TN = "{urn:oasis:names:tc:opendocument:xmlns:table:1.0}"
ON = "{urn:oasis:names:tc:opendocument:xmlns:office:1.0}"

CATEGORIES = [
    {"id": "all-crime", "label": "All recorded crime", "column": "Total recorded crime", "description": "All notifiable police-recorded offences excluding fraud. Anti-social behaviour incidents are not included."},
    {"id": "violence", "label": "Violence", "column": "Violence against the person", "description": "Violence against the person, including homicide, violence with and without injury, stalking and harassment. Recording practices affect comparisons."},
    {"id": "sexual-offences", "label": "Sexual offences", "column": "Sexual offences", "description": "Police-recorded rape and other sexual offences. Changes in reporting and recording can affect trends."},
    {"id": "robbery", "label": "Robbery", "column": "Robbery", "description": "Theft involving force or the threat of force against a person, including personal and business property. " + ROBBERY_RECORDING_NOTE},
    {"id": "burglary", "label": "Burglary", "column": "Burglary", "description": "Residential and non-residential burglary. Rates use resident population, rather than the number of households."},
    {"id": "vehicle-crime", "label": "Vehicle crime", "column": "Vehicle offences", "description": "Theft of and from vehicles, vehicle interference and related vehicle offences."},
    {"id": "theft-from-person", "label": "Theft from the person", "column": "Theft from the person", "description": "Theft directly from a person without force or threat of force, including pickpocketing."},
    {"id": "shoplifting", "label": "Shoplifting", "column": "Shoplifting", "description": "Police-recorded theft from shops. Counts depend on offences being reported to and recorded by police. From April 2025, clarified counting rules require shop theft involving violence or threats, including during escape, to be recorded as business robbery. This can affect year-on-year shoplifting comparisons spanning that date."},
    {"id": "criminal-damage", "label": "Criminal damage & arson", "column": "Criminal damage and arson", "description": "Recorded criminal damage to property and arson offences."},
    {"id": "drug-offences", "label": "Drug offences", "column": "Drug offences", "description": "Drug possession and trafficking offences. Recorded totals are influenced by police activity."},
    {"id": "weapons", "label": "Weapons possession", "column": "Possession of weapons offences", "description": "Possession of weapons offences. This is not a count of all knife-enabled crime."},
    {"id": "public-order", "label": "Public order", "column": "Public order offences", "description": "Police-recorded public order offences, including offences involving fear or provocation of violence."},
]
CAT_IDS = [c["id"] for c in CATEGORIES]
RECORDED_GROUPS = {
    "Violence against the person", "Sexual offences", "Robbery", "Theft offences",
    "Criminal damage and arson", "Drug offences", "Possession of weapons offences",
    "Public order offences", "Miscellaneous crimes against society",
}


def normalize(value: object) -> str:
    text = re.sub(r"\[note\s*\d+\]", "", str(value or ""), flags=re.I)
    return re.sub(r"\s+", " ", text).strip()


class Page(HTMLParser):
    def __init__(self, body: bytes):
        super().__init__(convert_charrefs=True)
        self.links: list[tuple[str, str]] = []
        self.text: list[str] = []
        self._href = None
        self._label: list[str] = []
        self.feed(body.decode("utf-8"))

    def handle_starttag(self, tag, attrs):
        if tag == "a":
            self._href = dict(attrs).get("href")
            self._label = []

    def handle_data(self, data):
        self.text.append(data)
        if self._href is not None:
            self._label.append(data)

    def handle_endtag(self, tag):
        if tag == "a" and self._href is not None:
            self.links.append((self._href, normalize(" ".join(self._label))))
            self._href = None

    def matching_link(self, predicate, base):
        for href, label in self.links:
            url = urllib.parse.urljoin(base, href)
            if predicate(urllib.parse.unquote(url), label):
                return url
        raise ValueError(f"Required source download link missing from {base}")

    def published_at(self):
        text = normalize(" ".join(self.text))
        found = re.search(r"Release date:\s*(\d{1,2}\s+[A-Za-z]+\s+\d{4})", text)
        if not found:
            raise ValueError("ONS publication date not found; review changed source page")
        return dt.datetime.strptime(found.group(1), "%d %B %Y").date().isoformat()


def download(url: str, *, limit: int = 180_000_000) -> bytes:
    """Verified HTTPS, bounded downloads, retry temporary network errors."""
    parsed = urllib.parse.urlparse(url)
    if parsed.scheme != "https" or parsed.hostname not in {
        "www.ons.gov.uk", "www.gov.uk", "assets.publishing.service.gov.uk",
    }:
        raise ValueError(f"Unexpected source host: {url}")
    for attempt in range(3):
        try:
            request = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0 (compatible; crime.compare official-data refresh)", "Accept": "*/*"})
            with urllib.request.urlopen(request, timeout=90) as response:
                data = response.read(limit + 1)
            if not data or len(data) > limit:
                raise ValueError(f"Empty or oversized response: {url}")
            return data
        except (urllib.error.URLError, TimeoutError, ConnectionError):
            if attempt == 2:
                raise
            time.sleep(attempt + 1)
    raise AssertionError("unreachable")


def source_file(source_id, title, page_url, file_url, published_at):
    body = download(file_url)
    if not zipfile.is_zipfile(io.BytesIO(body)):
        raise ValueError(f"Expected spreadsheet ZIP structure for {source_id}")
    digest = hashlib.sha256(body).hexdigest()
    path = CACHE / (digest + Path(urllib.parse.unquote(file_url)).suffix)
    CACHE.mkdir(parents=True, exist_ok=True)
    if not path.exists():
        path.write_bytes(body)
    record = {"id": source_id, "title": title, "url": page_url, "downloadUrl": file_url,
              "publishedAt": published_at, "sha256": digest, "bytes": len(body),
              "licence": "Open Government Licence v3.0"}
    return record, path


def ods_published_at(path):
    """Read the workbook's own update date, not a collection page's other releases."""
    with zipfile.ZipFile(path) as archive, archive.open("content.xml") as stream:
        for event, element in ET.iterparse(stream, events=("end",)):
            if element.tag == TN + "table-row":
                value = normalize(" ".join(element.itertext()))
                found = re.search(r"Updated:\s*(\d{1,2}\s+[A-Za-z]+\s+\d{4})", value)
                element.clear()
                if found:
                    return dt.datetime.strptime(found.group(1), "%d %B %Y").date().isoformat()
            if element.tag == TN + "table":
                break
    raise ValueError("No update date in Home Office workbook cover sheet")


def discover_sources():
    ons = Page(download(ONS_INDEX))
    ons_url = ons.matching_link(lambda u, _: "/policeforceareadatatables/" in u and u.endswith(".xlsx"), ONS_INDEX)
    ho = Page(download(HO_INDEX))
    ho_url = ho.matching_link(lambda u, label: u.endswith(".ods") and "Police Force Area" in label and "2013 onwards" in label, HO_INDEX)
    search = Page(download(POP_SEARCH))
    # Only accredited mid-year estimate tables, not similarly named experimental estimates.
    candidates = [urllib.parse.urljoin("https://www.ons.gov.uk", href) for href, _ in search.links
                  if "/populationestimates/adhocs/" in href
                  and "populationestimatesforpoliceforceareas" in href
                  and "bysingleyearofageandsex" in href]
    if not candidates:
        raise ValueError("No police-force population estimate publication found in ONS search")
    def population_end(url):
        years = re.findall(r"mid(20\d\d)", url)
        return max(map(int, years), default=0)
    pop_page_url = max(candidates, key=population_end)
    pop = Page(download(pop_page_url))
    pop_url = pop.matching_link(lambda u, _: u.endswith(".xlsx"), pop_page_url)
    specifications = [
        ("ons-crime", "ONS police force area crime tables", ONS_INDEX, ons_url, ons.published_at()),
        ("home-office-history", "Home Office revised quarterly recorded crime open data", HO_INDEX, ho_url, ons.published_at()),
        ("ons-population", "ONS police force area mid-year population estimates", pop_page_url, pop_url, pop.published_at()),
    ]
    records, paths = [], {}
    for spec in specifications:
        record, path = source_file(*spec)
        if record["id"] == "home-office-history":
            record["publishedAt"] = ods_published_at(path)
        records.append(record)
        paths[record["id"]] = path
        print(f"Fetched {record['id']}: {record['bytes']:,} bytes, {record['sha256'][:12]}")
    return records, paths


def spreadsheet_table(workbook, name):
    sheet = next((s for s in workbook if normalize(s.title) == name), None)
    if sheet is None:
        raise ValueError(f"Source worksheet {name} is missing")
    headers = None
    rows = {}
    region = None
    for row in sheet.iter_rows(values_only=True):
        if row[0] == "Area Code":
            headers = [normalize(c) for c in row]
            continue
        code = str(row[0] or "")
        if not re.fullmatch(r"[EKW]\d{8}", code):
            continue
        if headers is None:
            raise ValueError("Area table has no recognised header")
        if code.startswith("E120") or code == "W92000004":
            region = code
        name = normalize(row[1])
        if name == "WALES": name = "Wales"
        if name == "ENGLAND AND WALES": name = "England and Wales"
        if name == "ENGLAND": name = "England"
        if code == "E12000006": name = "East of England"
        rows[code] = {"name": name, "regionId": region, "values": dict(zip(headers, row))}
    return rows


def category_value(row, category):
    # Match exact labels, except the all-crime header carries the fraud note.
    if category["id"] == "all-crime":
        candidates = [v for h, v in row["values"].items() if h.startswith("Total recorded crime")]
    else:
        candidates = [v for h, v in row["values"].items() if h == category["column"]]
    if len(candidates) != 1:
        raise ValueError(f"Ambiguous/missing offence column: {category['column']}")
    value = candidates[0]
    if not isinstance(value, (int, float)) or isinstance(value, bool):
        return None
    if not math.isfinite(value) or value < 0 or value != int(value):
        raise ValueError(f"Invalid offence count {value}")
    return int(value)


def parse_crime(path):
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        current = spreadsheet_table(workbook, "Table P1")
        previous = spreadsheet_table(workbook, "Table P4")
        title = str(workbook["Table P1"].cell(1, 1).value)
        found = re.search(r"year ending (\w+) (20\d\d)", title, re.I)
        if not found:
            raise ValueError("Cannot determine crime reference period from ONS table")
        date = dt.datetime.strptime(found.group(1) + " " + found.group(2), "%B %Y").date()
        period_end = date.replace(day=calendar.monthrange(date.year, date.month)[1])
        if period_end.month not in (3, 6, 9, 12):
            raise ValueError("Expected a financial-quarter crime release")
        return current, previous, period_end
    finally:
        workbook.close()


def parse_population(path):
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    populations = collections.defaultdict(dict)
    try:
        recognised_sheets = 0
        for sheet in workbook:
            age_cols = None
            for row in sheet.iter_rows(values_only=True):
                if row[0] in ("PFA 2023 Code", "PFA Code") or (str(row[0]).startswith("PFA") and "Code" in str(row[0])):
                    # Sum every published female/male age band exactly once. No rounded totals.
                    age_cols = [i for i, value in enumerate(row) if re.fullmatch(r"[FM]\d+\+?", str(value or ""))]
                    if len(age_cols) < 180:
                        raise ValueError("Population age-band schema changed")
                    recognised_sheets += 1
                    continue
                if age_cols is None or not re.fullmatch(r"[EW]\d{8}", str(row[0] or "")):
                    continue
                found = re.fullmatch(r"mid-(20\d\d)", str(row[2] or ""))
                if not found:
                    raise ValueError("Population reference year schema changed")
                values = [row[i] for i in age_cols]
                if any(not isinstance(v, (int, float)) or v < 0 for v in values):
                    raise ValueError("Missing/invalid published population age cell")
                populations[row[0]][int(found.group(1))] = int(sum(values))
        if not recognised_sheets or len(populations) != 43:
            raise ValueError(f"Expected 43 territorial force populations, got {len(populations)}")
        return populations
    finally:
        workbook.close()


def quarter_end(financial_year, quarter):
    year = int(str(financial_year)[:4])
    month = {1: 6, 2: 9, 3: 12, 4: 3}[int(quarter)]
    if int(quarter) == 4:
        year += 1
    return dt.date(year, month, calendar.monthrange(year, month)[1])


def categories_for_offence(group, subgroup):
    if group not in RECORDED_GROUPS:
        return []
    result = ["all-crime"]
    mapping = {"Violence against the person": "violence", "Sexual offences": "sexual-offences",
               "Robbery": "robbery", "Criminal damage and arson": "criminal-damage",
               "Drug offences": "drug-offences", "Possession of weapons offences": "weapons",
               "Public order offences": "public-order"}
    if group in mapping:
        result.append(mapping[group])
    if group == "Theft offences":
        if "burglary" in subgroup.lower(): result.append("burglary")
        elif subgroup == "Vehicle offences": result.append("vehicle-crime")
        elif subgroup == "Theft from the person": result.append("theft-from-person")
        elif subgroup == "Shoplifting": result.append("shoplifting")
    return result


def parse_history(path, current, period_end):
    """Stream the 350MB uncompressed ODS; never materialise the full workbook."""
    name_to_code = {normalize(v["name"]).casefold(): code for code, v in current.items() if code.startswith(("E230", "W150"))}
    name_to_code["london, city of"] = CITY
    name_to_code["british transport police"] = "BTP"
    # Enough source quarters for 13 rolling-year points and the previous-year check.
    first_year = period_end.year - 5
    quarters = collections.defaultdict(lambda: collections.defaultdict(collections.Counter))
    sheet = ""
    table = None
    with zipfile.ZipFile(path) as archive, archive.open("content.xml") as stream:
        for event, element in ET.iterparse(stream, events=("start", "end")):
            if event == "start" and element.tag == TN + "table":
                sheet = element.get(TN + "name", "")
                table = element
            if event == "end" and element.tag == TN + "table-row":
                match = re.fullmatch(r"(20\d\d)_\d\d", sheet)
                if match and int(match.group(1)) >= first_year:
                    cells = []
                    for cell in element:
                        value = cell.get(ON + "value") or normalize(" ".join(cell.itertext()))
                        repeated = int(cell.get(TN + "number-columns-repeated", "1"))
                        cells.extend([value] * min(repeated, 8 - len(cells)))
                        if len(cells) >= 8: break
                    if len(cells) >= 8 and re.fullmatch(r"20\d\d/\d\d", str(cells[0])):
                        group, subgroup = cells[4], cells[5]
                        categories = categories_for_offence(group, subgroup)
                        if categories:
                            force = name_to_code.get(cells[2].casefold())
                            if force is None:
                                raise ValueError(f"Unknown territorial force in Home Office file: {cells[2]}")
                            date = quarter_end(cells[0], cells[1])
                            if date > period_end:
                                raise ValueError("Home Office and ONS latest periods disagree; wait for both releases")
                            count = float(cells[7])
                            if not count.is_integer() or count < 0:
                                raise ValueError("Invalid recent quarterly count")
                            for category in categories:
                                quarters[date][force][category] += int(count)
                element.clear()
                # Remove cleared row placeholders as well as their contents.
                if table is not None:
                    table.clear()
            if event == "end" and element.tag == TN + "table":
                element.clear()
                table = None
    dates = sorted(quarters)
    if not dates or dates[-1] != period_end:
        raise ValueError("Home Office history has no quarter matching ONS latest period")
    expected_forces = {code for code in current if code.startswith(("E230", "W150"))} | {"BTP"}
    for date in dates:
        if set(quarters[date]) != expected_forces:
            raise ValueError(f"Incomplete territorial-force coverage in {date}")
    rolling = collections.defaultdict(dict)
    for i in range(3, len(dates)):
        window = dates[i - 3:i + 1]
        # Explicit adjacency prevents accidentally treating a missing quarter as zero.
        indexes = [d.year * 4 + (d.month // 3) for d in window]
        if indexes != list(range(indexes[0], indexes[0] + 4)):
            raise ValueError("Gap in Home Office quarterly history")
        for code in expected_forces:
            rolling[dates[i]][code] = {cat: sum(quarters[d][code][cat] for d in window) for cat in CAT_IDS}
    return dict(sorted(rolling.items())[-13:])


def population_for(code, year, populations, members):
    codes = members.get(code, [code])
    values = []
    actual_years = []
    for force in codes:
        available = populations.get(force, {})
        valid = [y for y in available if y <= year]
        if not valid:
            raise ValueError(f"No actual resident-population denominator for {force} in {year}")
        actual = max(valid)
        if year - actual > 2:
            raise ValueError(f"Population source more than two years out of date for {force}")
        values.append(available[actual]); actual_years.append(actual)
    if len(set(actual_years)) != 1:
        raise ValueError("Mixed population estimate years in area aggregate")
    return sum(values), actual_years[0]


def build_dataset(records, paths):
    current, previous, period_end = parse_crime(paths["ons-crime"])
    populations = parse_population(paths["ons-population"])
    history = parse_history(paths["home-office-history"], current, period_end)
    force_codes = [code for code in current if code.startswith(("E230", "W150"))]
    if len(force_codes) != 43:
        raise ValueError("ONS current table does not contain all 43 territorial forces")
    members = collections.defaultdict(list)
    for code in force_codes:
        region = current[code]["regionId"]
        if region is None:
            raise ValueError(f"Cannot determine region for {code}")
        members[region].append(code)
    if len(members) != 10:
        raise ValueError("Expected nine English regions plus Wales")
    members[NATIONAL] = force_codes
    area_codes = [code for code in current if code in force_codes or code in members]
    # Ask for the reference year before the crime period end. population_for
    # chooses an actually published estimate and enforces a maximum age.
    latest_population_year = period_end.year - 1
    areas = []
    for code in area_codes:
        name = current[code]["name"]
        area_type = "force" if code in force_codes else "nation" if code == NATIONAL else "region"
        population, pop_year = population_for(code, latest_population_year, populations, members)
        metrics = {}
        for category in CATEGORIES:
            cat = category["id"]
            count = category_value(current[code], category)
            prev_count = category_value(previous[code], category)
            history_points = []
            for date, force_values in history.items():
                history_members = members.get(code, [code]) + (["BTP"] if code == NATIONAL else [])
                historic_count = sum(force_values[force][cat] for force in history_members)
                hist_pop, _ = population_for(code, min(date.year - 1, latest_population_year), populations, members)
                history_points.append({"periodEnd": date.isoformat(), "count": historic_count,
                                       "rate": None if code == CITY else round(historic_count / hist_pop * 1000, 4)})
            # Cross-check two independent official extracts. No invented fallback if
            # a new offence definition or file format makes them disagree.
            if count is not None and history_points[-1]["count"] != count:
                raise ValueError(f"Official source reconciliation failed: {name}/{cat}: ONS={count}, Home Office={history_points[-1]['count']}")
            last_year_date = period_end.replace(year=period_end.year - 1).isoformat()
            earlier = next((v["count"] for v in history_points if v["periodEnd"] == last_year_date), None)
            if prev_count is not None and earlier is not None and prev_count != earlier:
                raise ValueError(f"Previous-year reconciliation failed: {name}/{cat}: ONS={prev_count}, Home Office={earlier}")
            metrics[cat] = {"count": count,
                            "rate": None if count is None or code == CITY else round(count / population * 1000, 4),
                            "previousCount": prev_count,
                            "changePercent": None if count is None or not prev_count else round((count / prev_count - 1) * 100, 2),
                            "history": history_points}
        quality = []
        if code == CITY:
            quality.append("Resident-population rates are suppressed, following the ONS treatment. The City has a large visiting and commuter population and is excluded from rate rankings.")
        if area_type == "nation":
            quality.append("Includes British Transport Police. Regional and territorial-force totals exclude British Transport Police, whose crimes cannot be assigned consistently to these areas.")
        areas.append({"id": code, "code": code, "name": name, "type": area_type,
                      "regionId": current[code]["regionId"] if area_type == "force" else None,
                      "population": population, "populationYear": pop_year,
                      "comparisonEligible": code != CITY,
                      "metrics": metrics, "quality": quality})
    published_at = next(s["publishedAt"] for s in records if s["id"] == "ons-crime")
    period_start = dt.date(period_end.year - 1, period_end.month, 1)
    if period_end.month == 12:
        period_start = dt.date(period_end.year, 1, 1)
    else:
        period_start = period_start.replace(month=period_end.month + 1)
    generated = dt.datetime.now(dt.timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z")
    data = {"schemaVersion": 1, "generatedAt": generated, "scope": "England and Wales",
            "unit": "offences per 1,000 residents", "baselineId": NATIONAL,
            "coverage": {"start": period_start.isoformat(), "end": period_end.isoformat(),
                         "label": period_start.strftime("%b %Y") + " – " + period_end.strftime("%b %Y"),
                         "frequency": "quarterly", "sourcePublishedAt": published_at},
            "updatePolicy": {"checkFrequency": "monthly", "publicationFrequency": "quarterly",
                             "note": "Sources are checked each month. Comparable official crime figures are released quarterly, usually about four months after the period ends."},
            "categories": [{k: v for k, v in c.items() if k != "column"} for c in CATEGORIES],
            "sources": [{k: s[k] for k in ("id", "title", "url", "publishedAt")} for s in records],
            "notes": [
                "The England and Wales reference includes British Transport Police; regional and territorial-force comparisons exclude it. London regional totals include City of London.",
                ROBBERY_RECORDING_NOTE,
            ],
            "methodology": [
                "Counts describe police-recorded crime, rather than every crime that occurs. Reporting levels, recording practices and police activity affect comparisons and trends.",
                "Each count covers a rolling 12-month period. Year-on-year change compares the count with the preceding rolling year ending in the same quarter.",
                "Fraud and anti-social behaviour incidents are excluded. Weapons possession is not the same measure as knife-enabled crime.",
                "Rates are calculated from exact ONS mid-year resident-population estimates, not inferred from rounded published rates. Historical rates use the latest available estimate at or before the year preceding each period end. Historical population estimates may be revised.",
                "Crime histories sum four consecutive quarters from the latest Home Office revised open-data file. Latest and previous-year counts are reconciled against ONS Tables P1 and P4 before publication.",
                "City of London counts remain visible, but its resident-based rates and rate rankings are suppressed. London regional totals include City of London.",
                "This version covers England and Wales. Crime recording systems and offence definitions differ in Scotland and Northern Ireland.",
            ], "areas": areas}
    validate_dataset(data)
    manifest = {"schemaVersion": 1, "generatedAt": generated, "coverage": data["coverage"],
                "sources": records, "licence": "Open Government Licence v3.0",
                "methodologyReferences": [{"title": "ONS user guide: robbery definitions and April 2025 counting-rule clarification", "url": CLASSIFICATION_GUIDE}],
                "attribution": "Contains public sector information licensed under the Open Government Licence v3.0. Crime: Home Office and Office for National Statistics. Population: Office for National Statistics.",
                "transforms": ["ONS P1 current and P4 previous-year counts; Home Office four-quarter rolling sums; exact ONS age/sex population sums; offences divided by resident population × 1,000."],
                "validation": {"territorialForces": 43, "regionsIncludingWales": 10,
                               "categories": len(CATEGORIES), "rollingHistoryPoints": len(history),
                               "onsHomeOfficeReconciliation": "passed"}}
    return data, manifest


def validate_dataset(data):
    if data.get("schemaVersion") != 1:
        raise ValueError("Unsupported dataset schema")
    areas = data.get("areas", [])
    if len(areas) != 54 or len({a["id"] for a in areas}) != len(areas):
        raise ValueError("Expected 43 forces, ten regions, and one national baseline")
    types = collections.Counter(a["type"] for a in areas)
    if types != {"force": 43, "region": 10, "nation": 1}:
        raise ValueError("Area coverage mismatch")
    categories = {c["id"] for c in data["categories"]}
    if categories != set(CAT_IDS):
        raise ValueError("Unexpected category coverage")
    end = dt.date.fromisoformat(data["coverage"]["end"])
    start = dt.date.fromisoformat(data["coverage"]["start"])
    if (end - start).days not in (364, 365):
        raise ValueError("Latest comparison is not a complete rolling year")
    by_code = {area["id"]: area for area in areas}
    for area in areas:
        if not isinstance(area["population"], int) or area["population"] <= 0:
            raise ValueError("Missing actual population")
        if set(area["metrics"]) != categories:
            raise ValueError("Missing category metrics")
        if area["type"] == "force" and area["regionId"] not in by_code:
            raise ValueError("Missing regional geography mapping")
        for cat, metric in area["metrics"].items():
            for key in ("count", "previousCount"):
                value = metric[key]
                if value is not None and (not isinstance(value, int) or value < 0):
                    raise ValueError("Invalid published offence count")
            if metric["count"] is not None and metric["count"] > area["metrics"]["all-crime"]["count"]:
                raise ValueError("Category exceeds total crime")
            rate = metric["rate"]
            if area["id"] == CITY:
                if rate is not None or area["comparisonEligible"]:
                    raise ValueError("City of London rate must remain suppressed")
            elif metric["count"] is not None:
                expected_rate = metric["count"] / area["population"] * 1000
                if rate is None or abs(expected_rate - rate) > 0.000051:
                    raise ValueError("Incorrect population-normalised rate")
            if metric["count"] is not None and metric["previousCount"]:
                expected_change = (metric["count"] / metric["previousCount"] - 1) * 100
                if metric["changePercent"] is None or abs(expected_change - metric["changePercent"]) > 0.0051:
                    raise ValueError("Incorrect year-on-year change")
            history = metric["history"]
            if len(history) < 8 or history[-1]["periodEnd"] != data["coverage"]["end"]:
                raise ValueError("Incomplete trend history")
            if metric["count"] is not None and history[-1]["count"] != metric["count"]:
                raise ValueError("Trend endpoint differs from comparison count")
            if [p["periodEnd"] for p in history] != sorted({p["periodEnd"] for p in history}):
                raise ValueError("Unordered or duplicated history periods")
    for region in (a for a in areas if a["type"] == "region"):
        members = [a for a in areas if a["regionId"] == region["id"]]
        if not members or sum(a["population"] for a in members) != region["population"]:
            raise ValueError("Region population does not reconcile to its forces")
        for cat in categories:
            expected = sum(a["metrics"][cat]["count"] for a in members)
            if expected != region["metrics"][cat]["count"]:
                raise ValueError("Region offence count does not reconcile to its forces")
    # The national count includes BTP, so it is expected to exceed territorial sums.
    national = by_code[NATIONAL]
    territorial = [a for a in areas if a["type"] == "force"]
    if national["population"] != sum(a["population"] for a in territorial):
        raise ValueError("National resident-population aggregate does not reconcile")
    if national["metrics"]["all-crime"]["count"] < sum(a["metrics"]["all-crime"]["count"] for a in territorial):
        raise ValueError("Invalid national crime baseline")
    json.dumps(data, allow_nan=False)


def write_snapshot(data, manifest):
    OUTPUT.mkdir(parents=True, exist_ok=True)
    # Validate all output first, then replace complete files atomically. No truncation.
    payloads = [("crime.json", data), ("source-manifest.json", manifest)]
    temporary = []
    try:
        for filename, payload in payloads:
            with tempfile.NamedTemporaryFile(mode="w", encoding="utf-8", dir=OUTPUT, prefix=".refresh-", delete=False) as f:
                json.dump(payload, f, ensure_ascii=False, allow_nan=False, separators=(",", ":"))
                f.write("\n"); f.flush(); os.fsync(f.fileno())
                os.fchmod(f.fileno(), 0o644)
                temporary.append((Path(f.name), OUTPUT / filename))
        for source, target in temporary:
            os.replace(source, target)
    finally:
        for source, _ in temporary:
            source.unlink(missing_ok=True)


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--check", "--validate-only", dest="check", action="store_true", help="Validate the committed snapshot without fetching sources")
    parser.add_argument("--force", action="store_true", help="Rebuild even if all upstream file hashes are unchanged")
    args = parser.parse_args(argv)
    if args.check:
        validate_dataset(json.loads((OUTPUT / "crime.json").read_text()))
        manifest = json.loads((OUTPUT / "source-manifest.json").read_text())
        if len(manifest["sources"]) != 3 or manifest["validation"]["onsHomeOfficeReconciliation"] != "passed":
            raise ValueError("Invalid provenance manifest")
        print("Snapshot valid: 43 territorial forces, 10 regions, national baseline, 12 categories.")
        return 0
    records, paths = discover_sources()
    existing_manifest = OUTPUT / "source-manifest.json"
    if not args.force and existing_manifest.exists() and (OUTPUT / "crime.json").exists():
        old = json.loads(existing_manifest.read_text())
        if old.get("sources") == records:
            validate_dataset(json.loads((OUTPUT / "crime.json").read_text()))
            print("All official source files unchanged; published snapshot preserved.")
            return 0
    data, manifest = build_dataset(records, paths)
    # Refuse an upstream rollback even if its changed files parse successfully.
    if (OUTPUT / "crime.json").exists():
        existing = json.loads((OUTPUT / "crime.json").read_text())
        if data["coverage"]["end"] < existing["coverage"]["end"]:
            raise ValueError("Upstream data period rolled back; retaining newer published snapshot")
    write_snapshot(data, manifest)
    print(f"Published validated snapshot: {data['coverage']['label']}, {len(data['areas'])} areas.")
    return 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    except Exception as error:
        print(f"Refresh failed; existing published snapshot retained: {error}", file=sys.stderr)
        sys.exit(1)
